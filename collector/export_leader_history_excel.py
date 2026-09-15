"""
Excel de entrega de la Fase 1 (backfill histórico del líder). Seis pestañas:
Trader History, Live Capture, History Live Overlap, Data Quality,
Retrieval Manifest, Ambiguous Duplicates.

Fuentes: leader_trades_v2 (histórico canónico, nuevo) + leader_trades (poller
en vivo existente, sin tocar) + backfill_runs/backfill_pages (manifest) +
leader_trade_raw_links (trazabilidad de duplicados ambiguos). No toca mercados,
inventario, order book ni backtesting -- solo lee lo que ya existe.

Uso:  python3 export_leader_history_excel.py [ruta_salida.xlsx]
"""
import hashlib
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

import db
from config import LEADER_WALLET, REPO_ROOT

HISTORY_ORIGIN = "DATA_API_HISTORY"


def _fmt_ts(ts):
    if ts is None:
        return None
    return datetime.fromtimestamp(ts, timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _rows(conn, sql, *params):
    return [dict(r) for r in conn.execute(sql, params).fetchall()]


def _write_table(ws, headers, rows, start_row=1):
    """Escribe una tabla (headers + filas de dict) empezando en start_row.
    Devuelve la próxima fila libre."""
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    for i, row in enumerate(rows, start=start_row + 1):
        for j, h in enumerate(headers, start=1):
            ws.cell(row=i, column=j, value=row.get(h))
    return start_row + len(rows) + 1


def _sheet_trader_history(wb, conn, leader_wallet):
    ws = wb.create_sheet("Trader History")
    rows = _rows(conn, """
        SELECT id, transaction_hash, condition_id, market_slug, market_title, asset_symbol,
               token_id, outcome, side, price, shares, usdc_amount, source_timestamp_utc,
               multiplicity_index, multiplicity_total, dedup_ambiguous, run_id
        FROM leader_trades_v2
        WHERE origin=? AND leader_wallet=?
        ORDER BY source_timestamp_utc ASC
    """, HISTORY_ORIGIN, leader_wallet)
    for r in rows:
        r["source_timestamp_human_utc"] = _fmt_ts(r["source_timestamp_utc"])
    headers = ["id", "transaction_hash", "condition_id", "market_slug", "market_title",
               "asset_symbol", "token_id", "outcome", "side", "price", "shares", "usdc_amount",
               "source_timestamp_utc", "source_timestamp_human_utc", "multiplicity_index",
               "multiplicity_total", "dedup_ambiguous", "run_id"]
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A2"
    return rows


def _sheet_live_capture(wb, conn, leader_wallet):
    ws = wb.create_sheet("Live Capture")
    rows = _rows(conn, """
        SELECT id, transaction_hash, condition_id, market_slug, market_title, asset_symbol,
               token_id, outcome, side, price, shares, usdc_amount, source_timestamp_utc,
               detection_latency_ms, collection_method, is_startup_batch
        FROM leader_trades
        WHERE leader_wallet=? AND collection_method='LIVE' AND is_startup_batch=0
        ORDER BY source_timestamp_utc ASC
    """, leader_wallet)
    for r in rows:
        r["source_timestamp_human_utc"] = _fmt_ts(r["source_timestamp_utc"])
    headers = ["id", "transaction_hash", "condition_id", "market_slug", "market_title",
               "asset_symbol", "token_id", "outcome", "side", "price", "shares", "usdc_amount",
               "source_timestamp_utc", "source_timestamp_human_utc", "detection_latency_ms",
               "collection_method", "is_startup_batch"]
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A2"
    return rows


def _overlap_key(r):
    return (r.get("transaction_hash"), round(float(r["source_timestamp_utc"])),
            round(float(r["price"]), 6) if r.get("price") is not None else None,
            round(float(r["shares"]), 6) if r.get("shares") is not None else None,
            r.get("side"))


def _sheet_overlap(wb, history_rows, live_rows):
    ws = wb.create_sheet("History Live Overlap")
    hist_by_key = {}
    for r in history_rows:
        hist_by_key.setdefault(_overlap_key(r), []).append(r)
    live_by_key = {}
    for r in live_rows:
        live_by_key.setdefault(_overlap_key(r), []).append(r)

    hist_counts = Counter({k: len(v) for k, v in hist_by_key.items()})
    live_counts = Counter({k: len(v) for k, v in live_by_key.items()})

    out_rows = []
    for key in set(hist_counts) | set(live_counts):
        n_hist, n_live = hist_counts.get(key, 0), live_counts.get(key, 0)
        n_matched = min(n_hist, n_live)
        n_hist_only = n_hist - n_matched
        n_live_only = n_live - n_matched
        source_rows = hist_by_key.get(key) or live_by_key.get(key)
        base = source_rows[0]
        for status, n in (("matched", n_matched), ("history_only", n_hist_only), ("live_only", n_live_only)):
            for _ in range(n):
                out_rows.append({
                    "match_status": status, "transaction_hash": key[0],
                    "source_timestamp_utc": key[1], "price": key[2], "shares": key[3], "side": key[4],
                    "market_title": base.get("market_title"), "asset_symbol": base.get("asset_symbol"),
                })
    out_rows.sort(key=lambda r: (r["match_status"], r["source_timestamp_utc"] or 0))
    headers = ["match_status", "transaction_hash", "source_timestamp_utc", "price", "shares",
               "side", "market_title", "asset_symbol"]
    _write_table(ws, headers, out_rows)
    ws.freeze_panes = "A2"

    n_matched = sum(1 for r in out_rows if r["match_status"] == "matched")
    n_hist_only = sum(1 for r in out_rows if r["match_status"] == "history_only")
    n_live_only = sum(1 for r in out_rows if r["match_status"] == "live_only")
    return n_matched, n_hist_only, n_live_only


def _sheet_data_quality(wb, conn, leader_wallet, history_rows, live_rows, overlap_counts):
    ws = wb.create_sheet("Data Quality")
    n_matched, n_hist_only, n_live_only = overlap_counts

    total_raw = conn.execute(
        "SELECT count(*) c FROM leader_trades_raw WHERE origin=?", (HISTORY_ORIGIN,)
    ).fetchone()["c"]
    canonical_rows = len(history_rows)
    unique_keys = conn.execute(
        "SELECT count(DISTINCT natural_key) c FROM leader_trades_v2 WHERE origin=? AND leader_wallet=?",
        (HISTORY_ORIGIN, leader_wallet)).fetchone()["c"]
    exact_duplicate_observations = conn.execute("""
        SELECT count(*) c FROM leader_trade_raw_links l
        JOIN leader_trades_v2 v ON v.id = l.canonical_trade_id
        WHERE v.origin=? AND v.leader_wallet=? AND l.link_reason='repeat_observation'
    """, (HISTORY_ORIGIN, leader_wallet)).fetchone()["c"]
    ambiguous_groups = conn.execute("""
        SELECT count(DISTINCT natural_key) c FROM leader_trades_v2
        WHERE origin=? AND leader_wallet=? AND dedup_ambiguous=1
    """, (HISTORY_ORIGIN, leader_wallet)).fetchone()["c"]
    ambiguous_rows = sum(1 for r in history_rows if r["dedup_ambiguous"])

    buy_n = sum(1 for r in history_rows if r["side"] == "BUY")
    sell_n = sum(1 for r in history_rows if r["side"] == "SELL")
    other_side_n = canonical_rows - buy_n - sell_n

    by_asset = Counter(r["asset_symbol"] or "UNKNOWN" for r in history_rows)

    ts_values = [r["source_timestamp_utc"] for r in history_rows if r["source_timestamp_utc"] is not None]
    earliest, latest = (min(ts_values), max(ts_values)) if ts_values else (None, None)

    run = conn.execute("""
        SELECT * FROM backfill_runs WHERE origin=? AND leader_wallet=?
        ORDER BY started_at DESC LIMIT 1
    """, (HISTORY_ORIGIN, leader_wallet)).fetchone()

    pages_ok = conn.execute("""
        SELECT count(DISTINCT page_index) c, min(offset_requested) mn, max(offset_requested) mx
        FROM backfill_pages WHERE run_id=? AND error IS NULL
    """, (run["run_id"],)).fetchone() if run else {"c": 0, "mn": None, "mx": None}
    pages_failed_attempts = conn.execute("""
        SELECT count(*) c FROM backfill_pages WHERE run_id=? AND error IS NOT NULL
    """, (run["run_id"],)).fetchone()["c"] if run else 0

    if run and run["status"] == "completed_exhausted":
        verdict = f"PROBADA EXHAUSTIVA -- el lider no tiene mas historia (terminal_condition={run['terminal_condition']})"
    elif run and run["status"] == "completed_hard_limit":
        verdict = ("NO PROBADA EXHAUSTIVA -- BLOQUEADA POR LIMITE DOCUMENTADO DE LA API: "
                   f"{run['notes']}. La historia del lider puede continuar mas atras, pero "
                   "data-api/trades no la sirve por este canal.")
    elif run and run["status"] == "completed_safety_limit":
        verdict = "NO PROBADA -- se llego al safety_limit_pages sin condicion terminal real"
    elif run and run["status"] == "failed":
        verdict = f"NO PROBADA -- corrida detenida por {run['terminal_condition']} (probable falla transitoria)"
    elif run:
        verdict = f"EN CURSO -- status={run['status']}"
    else:
        verdict = "SIN CORRIDA REGISTRADA"

    metrics = [
        ("leader_wallet", leader_wallet),
        ("total_raw_observations", total_raw),
        ("canonical_trade_rows", canonical_rows),
        ("canonical_unique_natural_keys", unique_keys),
        ("exact_duplicate_raw_observations (solape/reintento, colapsados sin perderse)", exact_duplicate_observations),
        ("ambiguous_duplicate_groups (multiplicity_total>1)", ambiguous_groups),
        ("ambiguous_duplicate_canonical_rows", ambiguous_rows),
        ("buy_count", buy_n),
        ("sell_count", sell_n),
        ("other_side_count", other_side_n),
        ("earliest_timestamp_utc", earliest),
        ("earliest_timestamp_human_utc", _fmt_ts(earliest)),
        ("latest_timestamp_utc", latest),
        ("latest_timestamp_human_utc", _fmt_ts(latest)),
        ("pages_retrieved_ok", pages_ok["c"]),
        ("offset_min_retrieved", pages_ok["mn"]),
        ("offset_max_retrieved", pages_ok["mx"]),
        ("page_attempts_failed", pages_failed_attempts),
        ("run_id", run["run_id"] if run else None),
        ("run_status", run["status"] if run else None),
        ("run_terminal_condition", run["terminal_condition"] if run else None),
        ("TERMINAL_CONDITION_VERDICT", verdict),
        ("history_live_overlap_matched", n_matched),
        ("history_live_overlap_history_only", n_hist_only),
        ("history_live_overlap_live_only", n_live_only),
        ("generated_at_utc", _fmt_ts(time.time())),
    ]
    for asset, n in sorted(by_asset.items()):
        metrics.append((f"trades_by_asset:{asset}", n))

    ws.append(["metric", "value"])
    for k, v in metrics:
        ws.append([k, v])
    ws.freeze_panes = "A2"
    return verdict, run


def _sheet_retrieval_manifest(wb, conn, leader_wallet):
    ws = wb.create_sheet("Retrieval Manifest")
    runs = _rows(conn, "SELECT * FROM backfill_runs WHERE origin=? AND leader_wallet=? ORDER BY started_at",
                 HISTORY_ORIGIN, leader_wallet)
    for r in runs:
        r["started_at_human"] = _fmt_ts(r["started_at"])
        r["finished_at_human"] = _fmt_ts(r["finished_at"])
    run_headers = ["run_id", "origin", "leader_wallet", "started_at_human", "finished_at_human",
                   "status", "terminal_condition", "page_size", "safety_limit_pages",
                   "pages_requested", "pages_succeeded", "pages_failed", "last_offset_completed"]
    next_row = _write_table(ws, run_headers, runs)

    run_ids = [r["run_id"] for r in runs]
    pages = []
    for rid in run_ids:
        pages.extend(_rows(conn, """
            SELECT run_id, page_index, attempt_number, offset_requested, http_status,
                   cache_age_s, cache_status, n_rows_parsed, response_body_sha256,
                   normalized_page_fingerprint, error, request_started_at, response_received_at
            FROM backfill_pages WHERE run_id=? ORDER BY page_index, attempt_number
        """, rid))
    for r in pages:
        r["request_started_human"] = _fmt_ts(r["request_started_at"])
        r["response_received_human"] = _fmt_ts(r["response_received_at"])
    page_headers = ["run_id", "page_index", "attempt_number", "offset_requested", "http_status",
                    "cache_age_s", "cache_status", "n_rows_parsed", "response_body_sha256",
                    "normalized_page_fingerprint", "error", "request_started_human",
                    "response_received_human"]
    next_row += 1
    _write_table(ws, page_headers, pages, start_row=next_row)
    return runs, pages


def _sheet_ambiguous_duplicates(wb, conn, leader_wallet):
    ws = wb.create_sheet("Ambiguous Duplicates")
    rows = _rows(conn, """
        SELECT v.id canonical_trade_id, v.natural_key, v.multiplicity_index, v.multiplicity_total,
               v.transaction_hash, v.source_timestamp_utc, v.price, v.shares, v.side,
               v.market_title, l.link_reason, r.id raw_row_id, r.row_position_in_page,
               r.fetched_at, p.run_id, p.page_index, p.attempt_number
        FROM leader_trades_v2 v
        JOIN leader_trade_raw_links l ON l.canonical_trade_id = v.id
        JOIN leader_trades_raw r ON r.id = l.raw_row_id
        JOIN backfill_pages p ON p.id = r.page_id
        WHERE v.dedup_ambiguous=1 AND v.origin=? AND v.leader_wallet=?
        ORDER BY v.natural_key, v.multiplicity_index, l.link_reason DESC
    """, HISTORY_ORIGIN, leader_wallet)
    for r in rows:
        r["source_timestamp_human_utc"] = _fmt_ts(r["source_timestamp_utc"])
    headers = ["natural_key", "canonical_trade_id", "multiplicity_index", "multiplicity_total",
               "transaction_hash", "source_timestamp_utc", "source_timestamp_human_utc", "price",
               "shares", "side", "market_title", "link_reason", "raw_row_id",
               "row_position_in_page", "run_id", "page_index", "attempt_number"]
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A2"
    return rows


def build(out_path: Path, leader_wallet=LEADER_WALLET):
    db.init_db()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    with db.connect() as conn:
        history_rows = _sheet_trader_history(wb, conn, leader_wallet)
        live_rows = _sheet_live_capture(wb, conn, leader_wallet)
        overlap_counts = _sheet_overlap(wb, history_rows, live_rows)
        verdict, run = _sheet_data_quality(wb, conn, leader_wallet, history_rows, live_rows, overlap_counts)
        runs, pages = _sheet_retrieval_manifest(wb, conn, leader_wallet)
        ambiguous_rows = _sheet_ambiguous_duplicates(wb, conn, leader_wallet)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    body = out_path.read_bytes()
    sha256 = hashlib.sha256(body).hexdigest()
    (out_path.with_suffix(out_path.suffix + ".sha256")).write_text(f"{sha256}  {out_path.name}\n")

    n_matched, n_hist_only, n_live_only = overlap_counts
    summary = {
        "out_path": str(out_path),
        "sha256": sha256,
        "n_history_trades": len(history_rows),
        "n_live_trades": len(live_rows),
        "n_ambiguous_canonical_rows": len(ambiguous_rows),
        "history_earliest_utc": min((r["source_timestamp_utc"] for r in history_rows), default=None),
        "history_latest_utc": max((r["source_timestamp_utc"] for r in history_rows), default=None),
        "live_earliest_utc": min((r["source_timestamp_utc"] for r in live_rows), default=None),
        "live_latest_utc": max((r["source_timestamp_utc"] for r in live_rows), default=None),
        "terminal_condition_verdict": verdict,
        "run_id": run["run_id"] if run else None,
        "overlap_matched": n_matched, "overlap_history_only": n_hist_only, "overlap_live_only": n_live_only,
    }
    return summary


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO_ROOT / "export" / "leader_history.xlsx"
    summary = build(out)
    print("=" * 72)
    for k, v in summary.items():
        print(f"{k}: {v}")

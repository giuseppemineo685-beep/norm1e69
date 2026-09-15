"""
Excel final V1: patrones observados + Draft Strategy V1 + backtest cronologico
train/test sin look-ahead. Todo "estimated gross P&L" -- ningun numero acá es
P&L neto (sin fees, sin slippage mas alla del best_ask observado, sin riesgo
de ejecucion, sin redencion real). No es copy-trading: la regla se evalua
independientemente del lider en TODOS los mercados BTC/ETH/SOL de 5min.

Uso: python3 export_strategy_v1_excel.py --db-path /tmp/rebuild_test.db [salida.xlsx]
"""
import argparse
import csv
import hashlib
import sqlite3
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

import audit_strategy_v1 as aud
import backtest_momentum_v1 as bt


def fmt(ts):
    return None if ts is None else datetime.fromtimestamp(ts, timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _write_table(ws, headers, rows, start_row=1):
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    for i, row in enumerate(rows, start=start_row + 1):
        for j, h in enumerate(headers, start=1):
            v = row.get(h) if isinstance(row, dict) else row[j - 1]
            ws.cell(row=i, column=j, value=v)
    return start_row + len(rows) + 1


def momentum_pattern_analysis(strategy_csv):
    rows = list(csv.DictReader(open(strategy_csv)))

    def f(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    resolved = [r for r in rows if r["market_winner"] in ("Up", "Down")]
    buckets = [(240, 300), (180, 240), (120, 180), (60, 120), (30, 60), (0, 30)]
    out = []
    for lo, hi in buckets:
        a = c = aw = cw = 0
        for r in resolved:
            dist, sec, won = f(r["underlying_distance_from_open_pct"]), f(r["seconds_to_market_close"]), r["traded_side_won"]
            if dist is None or sec is None or won in (None, "") or abs(dist) < 0.01 or not (lo <= sec < hi):
                continue
            won = int(won)
            aligned = (dist > 0 and r["trade_outcome"] == "Up") or (dist < 0 and r["trade_outcome"] == "Down")
            if aligned:
                a += 1; aw += won
            else:
                c += 1; cw += won
        out.append({
            "seconds_remaining_bucket": f"[{lo}-{hi})", "n_aligned": a,
            "aligned_win_rate_pct": round(aw / a * 100, 1) if a else None,
            "n_contrarian": c, "contrarian_win_rate_pct": round(cw / c * 100, 1) if c else None,
        })

    n_total = len(resolved)
    n_up = sum(1 for r in rows if r["trade_outcome"] == "Up")
    n_down = sum(1 for r in rows if r["trade_outcome"] == "Down")
    return out, {"n_leader_trades_in_dataset": len(rows), "n_resolved": n_total, "n_up": n_up, "n_down": n_down}


def build(db_path, strategy_csv, out_path):
    result = bt.run(db_path)
    if result is None:
        raise SystemExit("backtest no produjo resultado -- ver salida de consola")

    print("\n" + "=" * 78 + "\nAUDITORIA FINAL (sin recalibrar)\n" + "=" * 78)
    audit = aud.run(db_path)

    pattern_rows, pattern_meta = momentum_pattern_analysis(strategy_csv)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------- Overview ----------
    ws = wb.create_sheet("Overview")
    ws["A1"] = "V1 exploratoria -- NO copy-trading. La estrategia decide sola, mirando el mercado, no al lider."
    ws["A2"] = "TODOS los P&L de este archivo son ESTIMATED GROSS P&L: sin fees, sin slippage mas alla del " \
               "best_ask observado en el snapshot, sin riesgo de ejecucion real, sin redencion on-chain real."
    ws["A3"] = ("HALLAZGO DE LA AUDITORIA: el baseline 'comprar el favorito por precio' (83 trades, mismo "
                "universo TEST, sin ningun umbral) obtiene ROI 8.35% y 77.1% win rate -- igual o mejor que "
                "V1 (63 trades, ROI 7.15%, 76.2%). El 100% de los trades de V1 son un subconjunto de los "
                "del favorito, y coinciden de lado en el 85.7% de los casos. V1 NO demuestra una ventaja "
                "clara sobre simplemente comprar lo que el mercado ya cree mas probable. Ver Baselines Comparison.")
    overview = [
        ("generado_utc", fmt(time.time())),
        ("db_usada", db_path),
        ("objetivo", "Dataset de decisiones del lider -> patrones de mercado -> regla autonoma -> backtest"),
        ("trades_del_lider_en_dataset_limpio", pattern_meta["n_leader_trades_in_dataset"]),
        ("de_esos_resueltos", pattern_meta["n_resolved"]),
        ("regla_v1", f"momentum de apertura, decision a open+{bt.DECISION_OFFSET_S}s, "
                     f"umbral elegido en TRAIN={result['chosen_threshold_pct']}%"),
        ("mercados_candidatos_totales", result["n_candidates"]),
        ("descartados_sin_subyacente_en_el_instante", result["n_no_underlying"]),
        ("descartados_sin_order_book_del_lado_señalado", result["n_no_book"]),
        ("train_rango", " -> ".join(result["train_range"]) if result["train_range"] else None),
        ("test_rango", " -> ".join(result["test_range"]) if result["test_range"] else None),
        ("TEST_n_trades", result["test_result"]["n_trades"]),
        ("TEST_win_rate_pct", round(result["test_result"]["win_rate"] * 100, 1) if result["test_result"]["win_rate"] else None),
        ("TEST_estimated_gross_pnl_usd", round(result["test_result"]["gross_pnl"], 2)),
        ("TEST_roi_estimated_pct", round(result["test_result"]["roi"] * 100, 2) if result["test_result"]["roi"] else None),
        ("TEST_con_precio_ejecutable_real__fill_status", f"FULL={audit['n_full']} PARTIAL={audit['n_partial']} SKIPPED={audit['n_skipped']}"),
        ("TEST_ejecutable__usd_desplegado_real", round(audit["usd_deployed_total"], 2)),
        ("TEST_ejecutable__gross_pnl_usd", round(audit["pnl_exec_total"], 2)),
        ("TEST_ejecutable__win_rate_pct", round(audit["n_wins_exec"] / len(audit["executable_trades"]) * 100, 1) if audit["executable_trades"] else None),
        ("baseline_favorito__n_trades_pnl_roi", f"{audit['baseline_favorite']['n_trades']} / "
         f"${audit['baseline_favorite'].get('gross_pnl_usd')} / {audit['baseline_favorite'].get('roi_pct')}%"),
        ("baseline_sin_umbral__n_trades_pnl_roi", f"{audit['baseline_momentum_raw']['n_trades']} / "
         f"${audit['baseline_momentum_raw'].get('gross_pnl_usd')} / {audit['baseline_momentum_raw'].get('roi_pct')}%"),
        ("baseline_aleatorio__n_trades_pnl_roi", f"{audit['baseline_random']['n_trades']} / "
         f"${audit['baseline_random'].get('gross_pnl_usd')} / {audit['baseline_random'].get('roi_pct')}% "
         f"(seed={audit['random_seed']})"),
        ("anti_lookahead_subyacente_violaciones", audit["viol_under"]),
        ("anti_lookahead_snapshot_violaciones", audit["viol_snap"]),
    ]
    ws.append(["campo", "valor"])
    for k, v in overview:
        ws.append([k, v])
    ws.freeze_panes = "A5"

    # ---------- Trader Pattern Analysis ----------
    ws = wb.create_sheet("Trader Pattern Analysis")
    ws["A1"] = ("Patron descubierto en el dataset limpio del lider (usable_for_strategy_learning=1): "
                "cuando el lado que compra coincide con el signo de underlying_distance_from_open_pct "
                "en ese instante ('aligned'), gana mucho mas seguido que cuando va en contra ('contrarian'). "
                "Se muestra por segundos restantes para descartar que sea un artefacto de trades tarde en la ventana.")
    _write_table(ws, ["seconds_remaining_bucket", "n_aligned", "aligned_win_rate_pct",
                       "n_contrarian", "contrarian_win_rate_pct"], pattern_rows, start_row=3)
    ws.freeze_panes = "A4"

    # ---------- Draft Strategy V1 ----------
    ws = wb.create_sheet("Draft Strategy V1")
    rules = [
        ("nombre", "momentum_apertura_v1"),
        ("tipo", "Autonoma -- NO copia al lider, no observa sus trades en tiempo real"),
        ("universo", "Mercados Up/Down de 5 minutos, BTC/ETH/SOL"),
        ("instante_de_decision", f"open_time_utc + {bt.DECISION_OFFSET_S}s (una sola vez por mercado)"),
        ("señal", "signo(underlying_distance_from_open_pct) en el instante de decision"),
        ("filtro_de_entrada", f"|underlying_distance_from_open_pct| >= {result['chosen_threshold_pct']}% "
                              "(umbral elegido en TRAIN, nunca reajustado con TEST)"),
        ("lado", "Up si distancia>0, Down si distancia<0. Sin señal (distancia~0): NO TRADE"),
        ("tamaño", f"stake fijo ${bt.STAKE_USD} por trade"),
        ("precio_de_entrada", "best_ask del snapshot mas reciente <= instante de decision (misma regla "
                              "anti-lookahead que trade_context: event_ts Y received_ts <= decision_ts)"),
        ("salida", "ninguna -- se mantiene a resolucion del mercado (payout $1 si gana, $0 si pierde)"),
        ("criterio_de_seleccion_del_umbral", "el que maximiza gross_pnl en TRAIN entre los umbrales con >=20 trades"),
        ("NO_incluye", "fees, rebates, slippage mas alla del best_ask observado, fill parcial, "
                       "latencia de orden propia, riesgo de que el ask se mueva entre decision y ejecucion real"),
    ]
    ws.append(["parametro", "valor"])
    for k, v in rules:
        ws.append([k, v])
    ws.freeze_panes = "A2"

    # ---------- Backtest Train ----------
    ws = wb.create_sheet("Backtest Train")
    ws["A1"] = "Barrido de umbral SOLO sobre TRAIN -- usado para elegir el umbral, nunca para reportar performance final."
    _write_table(ws, ["threshold_pct", "n_trades", "n_wins", "n_losses", "win_rate", "stake_total",
                       "payoff_total", "gross_pnl", "roi"], result["train_sweep"], start_row=3)
    ws.freeze_panes = "A4"

    # ---------- Backtest Test ----------
    ws = wb.create_sheet("Backtest Test")
    ws["A1"] = ("Los 63 trades TEST, individualmente. Dos modelos de precio de entrada uno al lado del otro: "
                "'simple' (best_ask, asume fill completo -- el de la corrida anterior) y 'executable' "
                "(recorre profundidad real para llenar $10; FULL/PARTIAL/SKIPPED, nunca asume fill).")
    headers = ["condition_id", "market_title", "decision_time_utc", "asset_symbol",
               "distance_from_open_pct", "signal_side", "best_ask", "depth_available_usd",
               "fill_status", "shares_simulated_simple", "shares_filled_executable",
               "usd_deployed_executable", "vwap_executable", "winner", "won",
               "pnl_simple_usd", "pnl_executable_usd"]
    _write_table(ws, headers, audit["detail_rows"], start_row=3)
    ws.freeze_panes = "A4"

    # ---------- Exclusion Funnel ----------
    ws = wb.create_sheet("Exclusion Funnel")
    ws["A1"] = (f"Universo TEST completo (open_time_utc >= {fmt(audit['test_boundary_ts'])}, BTC/ETH/SOL "
                f"5min resueltos): {audit['n_all_test_markets']} mercados. Por que cada uno NO termino "
                f"siendo un trade ejecutado de V1 (o si lo fue).")
    funnel_order = ["hueco_vpn", "contexto_incompleto_subyacente", "contexto_incompleto_orderbook",
                     "sin_señal_distancia_cero", "sin_señal_bajo_umbral", "ventana_invalida",
                     "liquidez_insuficiente_sin_profundidad", "señal_ejecutada_fill_parcial",
                     "señal_ejecutada_fill_completo"]
    funnel_labels = {
        "hueco_vpn": "Excluido: hueco de order book (VPN caida, 14:58:30-16:35:18 UTC)",
        "contexto_incompleto_subyacente": "Excluido: sin lectura de subyacente en el instante de decision",
        "contexto_incompleto_orderbook": "Excluido: sin order book del lado señalado en el instante de decision",
        "sin_señal_distancia_cero": "Excluido: sin señal (distancia del subyacente exactamente 0)",
        "sin_señal_bajo_umbral": f"Excluido: señal por debajo del umbral ({result['chosen_threshold_pct']}%)",
        "ventana_invalida": "Excluido: ventana de mercado invalida (decision cae fuera de la ventana)",
        "liquidez_insuficiente_sin_profundidad": "Señal SI, pero SKIPPED por falta de profundidad real (solo top-of-book)",
        "señal_ejecutada_fill_parcial": "Ejecutado con fill PARCIAL (profundidad insuficiente para $10 completos)",
        "señal_ejecutada_fill_completo": "Ejecutado con fill COMPLETO -- estos son los trades reales de V1",
    }
    funnel_rows = [{"categoria": funnel_labels[k], "n_mercados": audit["exclusion_funnel"].get(k, 0)}
                    for k in funnel_order]
    funnel_rows.append({"categoria": "TOTAL", "n_mercados": sum(audit["exclusion_funnel"].values())})
    _write_table(ws, ["categoria", "n_mercados"], funnel_rows, start_row=3)
    ws.column_dimensions["A"].width = 75
    ws.freeze_panes = "A4"

    # ---------- Baselines Comparison ----------
    ws = wb.create_sheet("Baselines Comparison")
    ws["A1"] = ("Las 4 estrategias evaluadas sobre el MISMO universo de 83 candidatos TEST (subyacente + "
                "order book disponibles en el instante de decision). Modelo de entrada 'simple' (best_ask, "
                "fill asumido completo) para las 4, comparacion apples-to-apples. Aleatorio con seed fija "
                f"({audit['random_seed']}), reproducible.")
    baseline_rows = []
    for ev in (audit["v1_eval"], audit["baseline_favorite"], audit["baseline_momentum_raw"], audit["baseline_random"]):
        row = dict(ev)
        exp = row.pop("exposure_by_asset", {})
        for asset in ("BTC", "ETH", "SOL"):
            row[f"n_trades_{asset}"] = exp.get(asset, 0)
        baseline_rows.append(row)
    headers2 = ["strategy", "n_trades", "n_wins", "win_rate_pct", "stake_total_usd", "gross_pnl_usd",
                "roi_pct", "max_drawdown_usd", "n_trades_BTC", "n_trades_ETH", "n_trades_SOL"]
    _write_table(ws, headers2, baseline_rows, start_row=3)
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "A4"

    # ---------- Results Summary ----------
    ws = wb.create_sheet("Results Summary")
    tr = result["test_result"]
    trr = result["train_at_chosen"]
    summary = [
        ("periodo", "TRAIN (calibracion)", "TEST (validacion, nunca tocado durante calibracion)"),
        ("rango_temporal", " -> ".join(result["train_range"]) if result["train_range"] else "-",
         " -> ".join(result["test_range"]) if result["test_range"] else "-"),
        ("n_trades", trr["n_trades"], tr["n_trades"]),
        ("win_rate", f"{trr['win_rate']*100:.1f}%" if trr["win_rate"] else "-",
         f"{tr['win_rate']*100:.1f}%" if tr["win_rate"] else "-"),
        ("stake_total_usd", trr["stake_total"], tr["stake_total"]),
        ("estimated_gross_pnl_usd", round(trr["gross_pnl"], 2), round(tr["gross_pnl"], 2)),
        ("roi_estimated", f"{trr['roi']*100:.2f}%" if trr["roi"] else "-",
         f"{tr['roi']*100:.2f}%" if tr["roi"] else "-"),
    ]
    for row in summary:
        ws.append(list(row))
    ws.freeze_panes = "A2"

    # ---------- Limitations ----------
    ws = wb.create_sheet("Limitations")
    limitations = [
        "Un solo dia de datos (2026-09-15, ~16h). Ningun resultado acá es una conclusion estadistica "
        "robusta -- es exploratorio, y hay que repetirlo sobre mas dias antes de confiar en el.",
        "Sesgo de cobertura del collector: de 588 mercados BTC/ETH/SOL de 5min resueltos, solo 206 "
        "(35%) tenian subyacente Y order book disponibles en el instante de decision -- los candidatos "
        "usados NO son una muestra aleatoria, son los mercados donde el collector capturaba bien.",
        "El umbral se eligio maximizando P&L bruto en TRAIN entre candidatos con >=20 trades -- es una "
        "decision editorial, no la unica razonable (optimizar ROI o win_rate habria elegido otro umbral).",
        "P&L BRUTO (estimated gross P&L): no incluye fees, rebates, ni redencion real on-chain.",
        "El precio de entrada es el best_ask observado en el snapshot mas cercano -- no modela que ese "
        "ask pueda moverse o desaparecer entre la decision y una orden real, ni fill parcial.",
        "Tamaño de posicion fijo ($10/trade) -- no hay gestion de riesgo, sizing dinamico ni limite de "
        "exposicion simultanea entre mercados solapados.",
        "El patron de 'momentum de apertura' se descubrio MIRANDO los trades del lider como muestreo -- "
        "el lider elige QUE momentos mirar, así que el patron podria reflejar parcialmente su propio "
        "buen ojo para elegir momentos, no una propiedad universal del mercado. El backtest sí evalua "
        "la regla en TODOS los mercados candidatos, no solo los que el lider opero -- pero la MUESTRA "
        "de candidatos con datos completos igual esta influenciada por cuando el collector cubria bien "
        "(que correlaciona con cuando el lider estaba activo).",
        "121 trades del hueco de order book (14:58:30-16:35:18 UTC) excluidos explicitamente de todo "
        "este analisis, tal como se pidio.",
        "Ningun snapshot recibido despues del trade se uso en ningun calculo (regla conservadora "
        "event_ts Y received_ts <= instante de decision, verificada con 0 violaciones).",
        "underlying = Binance spot, no la fuente exacta (Chainlink) que usa Polymarket para resolver "
        "-- documentado como aproximacion, no validado como identico.",
        "CRITICO (auditoria final): el baseline 'comprar el favorito por precio de Polymarket' (83 "
        "trades, sin ningun umbral ni logica de momentum) iguala o supera a V1 en el mismo universo TEST "
        "(ROI 8.35% vs 7.15%, win rate 77.1% vs 76.2%). El 100% de los trades de V1 son un subconjunto "
        "de los del favorito y coinciden de lado el 85.7% de las veces. V1 NO demuestra una ventaja "
        "incremental clara sobre simplemente comprar lo que el precio de mercado ya implica como mas "
        "probable -- el 'patron de momentum' descubierto podria ser, en gran parte, el propio mercado "
        "ya siendo eficiente, no una ineficiencia explotable nueva.",
        "Con precio ejecutable real (recorriendo profundidad, no asumiendo fill): 6 de los 63 trades de "
        "V1 quedan SKIPPED por falta de profundidad registrada mas alla del top-of-book (no hay niveles "
        "guardados en ese snapshot) -- ese P&L simplemente no se cuenta, no se le asume ni exito ni "
        "fracaso. De los 57 restantes, el fill fue completo en todos (ninguno PARCIAL en esta muestra).",
    ]
    ws.append(["limitacion"])
    for l in limitations:
        ws.append([l])
    ws.column_dimensions["A"].width = 120

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    sha256 = hashlib.sha256(out_path.read_bytes()).hexdigest()
    out_path.with_suffix(out_path.suffix + ".sha256").write_text(f"{sha256}  {out_path.name}\n")
    print(f"wrote {out_path}  sha256={sha256}")
    return out_path, sha256, result


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--db-path", required=True)
    p.add_argument("--strategy-csv", default=str(Path(__file__).resolve().parent / "export" / "strategy_learning_dataset.csv"))
    p.add_argument("out_xlsx", nargs="?", default=str(Path(__file__).resolve().parent / "export" / "strategy_v1.xlsx"))
    args = p.parse_args()
    build(args.db_path, args.strategy_csv, args.out_xlsx)

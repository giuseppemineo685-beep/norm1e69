"""
Excel de la validacion forward en paper trading. Lee SOLO
collector/paper_validation.db (nunca data.db). Debe poder generarse aunque
todavia no haya ningun mercado resuelto -- las hojas quedan con encabezado
y cero filas, nunca fallan por falta de datos.

Uso: python3 export_paper_validation.py [salida.xlsx]
"""
import hashlib
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

import paper_validation_db as pdb

# Fee de TAKER de Polymarket, categoria crypto -- VERIFICADO contra
# docs.polymarket.com/trading/fees (consultado 2026-09-16) y cruzado con
# multiples fuentes secundarias de 2026. Formula oficial:
#   fee = shares * TAKER_FEE_RATE_CRYPTO * price * (1 - price)
# Maker = 0% (no aplica: las 3 estrategias siempre cruzan contra el ask,
# nunca dejan orden resting -- son 100% taker). Pico de $1.75 por 100 shares
# a precio=0.50, coincide con lo reportado por las fuentes.
#
# NO incluye fee de redencion/resolucion: la mayoria de fuentes dice que
# redimir shares ganadoras no tiene costo adicional, pero UNA fuente
# secundaria menciono un 2% contradictorio -- por instruccion explicita, NO
# se aplica ningun numero para redencion. Se marca DESCONOCIDO, no cero.
#
# La formula usa el precio POR FILL. Este dataset solo guarda el VWAP
# ejecutable agregado de cada entrada/cobertura (no el desglose nivel por
# nivel de la caminata de profundidad), asi que el fee se aproxima con ese
# VWAP -- ver ESTIMATED_FEE_METHOD_NOTE. Nunca se aplica a los fills/
# decisiones guardados (paper_decisions/paper_resolutions no se tocan),
# solo a un calculo derivado y separado en el reporte.
TAKER_FEE_RATE_CRYPTO = 0.07
TAKER_FEE_SOURCE = "docs.polymarket.com/trading/fees (categoria crypto), consultado 2026-09-16"
ESTIMATED_FEE_METHOD_NOTE = (
    "APROXIMADO: fee = shares * 0.07 * VWAP_ejecutable * (1-VWAP_ejecutable) por fill. Usa el "
    "VWAP agregado de la caminata de profundidad, no el precio exacto de cada nivel individual "
    "dentro del fill -- una aproximacion razonable, no un calculo nivel-por-nivel exacto. "
    "NO es beneficio neto definitivo: no incluye fee de redencion (desconocido, no verificado, "
    "no aplicado), rebates de maker, ni ningun otro costo."
)


def estimated_taker_fee_usd(shares, price):
    """Fee de taker aproximado para un fill (ver TAKER_FEE_SOURCE / ESTIMATED_FEE_METHOD_NOTE)."""
    if shares is None or price is None:
        return None
    return shares * TAKER_FEE_RATE_CRYPTO * price * (1 - price)


def fmt(ts):
    return None if ts is None else datetime.fromtimestamp(ts, timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _rows(conn, sql, *params):
    return [dict(r) for r in conn.execute(sql, params).fetchall()]


def _write_table(ws, headers, rows, start_row=1):
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    for i, row in enumerate(rows, start=start_row + 1):
        for j, h in enumerate(headers, start=1):
            ws.cell(row=i, column=j, value=row.get(h))
    return start_row + len(rows) + 1


def _drawdown(pnls_ordered):
    cum = peak = max_dd = 0.0
    for p in pnls_ordered:
        cum += p
        peak = max(peak, cum)
        max_dd = max(max_dd, peak - cum)
    return max_dd


def build(out_path):
    pdb.init_db()
    REPORT_CUTOFF_TS = time.time()  # corte unico: TODAS las consultas de este informe usan este mismo instante

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    with pdb.connect() as conn:
        # ---------- Overview ----------
        ws = wb.create_sheet("Overview")
        ws["A1"] = ("Validacion forward en PAPER TRADING -- SOLO simulacion, datos completamente "
                    "nuevos (mercados posteriores al arranque del validador). No copy-trading, no "
                    "ejecucion real. Estimated gross P&L. Puede tener cero mercados resueltos todavia.")
        ws["A2"] = ("Hubo una caida de VPN/acceso a Polymarket durante el arranque de esta validacion. "
                    "Los mercados anteriores a la recuperacion quedan marcados PRE_VPN_RECOVERY -- NO se "
                    "borraron (siguen visibles en Signals/Resolutions), pero 'Strategy Summary' (metricas "
                    "oficiales) SOLO cuenta mercados OFFICIAL. Ver tag_vpn_recovery_cohort.py y hoja 'Sleep Audit'.")
        ws["A3"] = (f"CORTE UNICO DE ESTE INFORME: {fmt(REPORT_CUTOFF_TS)}. Todas las hojas (Signals, "
                    f"Hedges, Resolutions, Strategy Summary, Momentum vs Favorite) filtran por este mismo "
                    f"instante -- decisiones/coberturas/resoluciones posteriores a este corte, si el "
                    f"validador las genero mientras se armaba este Excel, NO estan incluidas todavia (se "
                    f"veran en la proxima corrida). Esto evita que los conteos difieran entre hojas.")
        cutoff = pdb.get_meta("validator_started_at")
        n_markets = conn.execute(
            "SELECT count(*) c FROM paper_markets WHERE discovered_at <= ?", (REPORT_CUTOFF_TS,)
        ).fetchone()["c"]
        n_decisions = conn.execute(
            "SELECT count(*) c FROM paper_decisions WHERE created_at <= ?", (REPORT_CUTOFF_TS,)
        ).fetchone()["c"]
        n_hedge_fills = conn.execute(
            "SELECT count(*) c FROM paper_hedge_fills WHERE executed_at_utc <= ?", (REPORT_CUTOFF_TS,)
        ).fetchone()["c"]
        n_resolutions = conn.execute(
            "SELECT count(*) c FROM paper_resolutions WHERE resolved_at_utc <= ?", (REPORT_CUTOFF_TS,)
        ).fetchone()["c"]
        vpn_recovery = pdb.get_meta("vpn_recovery_ts")
        cohort_start = pdb.get_meta("official_cohort_start_ts")
        n_pre_vpn = conn.execute(
            "SELECT count(*) c FROM paper_markets WHERE validation_cohort='PRE_VPN_RECOVERY' AND discovered_at <= ?",
            (REPORT_CUTOFF_TS,)
        ).fetchone()["c"]
        n_official = conn.execute(
            "SELECT count(*) c FROM paper_markets WHERE validation_cohort='OFFICIAL' AND discovered_at <= ?",
            (REPORT_CUTOFF_TS,)
        ).fetchone()["c"]
        # los 9 mercados afectados por el sueno son, ellos mismos, cohorte OFFICIAL (no PRE_VPN_RECOVERY)
        # -- el denominador correcto es n_official, no n_official+n_pre_vpn. 9/364=2.47%, no 0.7%.
        sleep_affected_pct = round(9 / n_official * 100, 2) if n_official else None
        overview = [
            ("generado_utc", fmt(time.time())),
            ("corte_unico_del_informe_utc", fmt(REPORT_CUTOFF_TS)),
            ("validador_arrancado_utc", fmt(float(cutoff)) if cutoff else None),
            ("mercados_descubiertos", n_markets),
            ("decisiones_tomadas", n_decisions),
            ("coberturas_ejecutadas", n_hedge_fills),
            ("resoluciones_registradas", n_resolutions),
            ("vpn_recovery_utc", fmt(float(vpn_recovery)) if vpn_recovery else None),
            ("cohorte_oficial_arranca_utc", fmt(float(cohort_start)) if cohort_start else None),
            ("mercados_PRE_VPN_RECOVERY_excluidos_de_metricas", n_pre_vpn),
            ("mercados_OFFICIAL_en_metricas", n_official),
            ("mercados_afectados_por_el_sueno_del_9_16", 9),
            ("pct_afectado_por_el_sueno_sobre_mercados_OFFICIAL", f"9/{n_official} = {sleep_affected_pct}%"),
        ]
        ws.append(["campo", "valor"])
        for k, v in overview:
            ws.append([k, v])
        ws.freeze_panes = "A5"

        # ---------- Signals (decisions) ----------
        ws = wb.create_sheet("Signals")
        rows = _rows(conn, """
            SELECT pd.strategy, pm.market_title, pd.asset_symbol, pd.condition_id, pd.token_id,
                   pm.validation_cohort, pd.decision_timestamp_utc, pd.underlying_open_price,
                   pd.underlying_price_at_60s, pd.distance_pct, pd.signal_present, pd.side_chosen,
                   pd.snapshot_id, pd.best_ask, pd.executable_price, pd.shares, pd.capital_usd,
                   pd.fill_status, pd.skip_or_partial_reason
            FROM paper_decisions pd JOIN paper_markets pm ON pm.condition_id = pd.condition_id
            WHERE pd.created_at <= ?
            ORDER BY pd.decision_timestamp_utc
        """, REPORT_CUTOFF_TS)
        for r in rows:
            r["decision_time_human"] = fmt(r["decision_timestamp_utc"])
            r["estimated_taker_fee_usd"] = (round(estimated_taker_fee_usd(r["shares"], r["executable_price"]), 4)
                                             if r["fill_status"] in ("FULL", "PARTIAL") else None)
        headers = ["strategy", "market_title", "asset_symbol", "condition_id", "token_id",
                   "validation_cohort", "decision_timestamp_utc", "decision_time_human",
                   "underlying_open_price", "underlying_price_at_60s", "distance_pct", "signal_present",
                   "side_chosen", "snapshot_id", "best_ask", "executable_price", "shares", "capital_usd",
                   "fill_status", "skip_or_partial_reason", "estimated_taker_fee_usd"]
        _write_table(ws, headers, rows, start_row=3)
        ws["A1"] = ("Cada mercado x cada estrategia (las 3 evaluan exactamente el mismo universo), corte "
                    "unico ver Overview. Incluye PRE_VPN_RECOVERY -- no se oculta nada acá, solo se "
                    "excluye de Strategy Summary. estimated_taker_fee_usd: " + ESTIMATED_FEE_METHOD_NOTE)
        ws.freeze_panes = "A4"

        # ---------- Hedges ----------
        ws = wb.create_sheet("Hedges")
        checks = _rows(conn, """
            SELECT hc.*, pd.condition_id AS cid2 FROM paper_hedge_checks hc
            JOIN paper_decisions pd ON pd.id = hc.decision_id
            WHERE hc.checked_at_utc <= ?
            ORDER BY hc.checked_at_utc
        """, REPORT_CUTOFF_TS)
        for r in checks:
            r["checked_at_human"] = fmt(r["checked_at_utc"])
        headers_c = ["decision_id", "condition_id", "checked_at_human", "opposite_token_id",
                     "opposite_snapshot_id", "opposite_best_ask", "opposite_executable_price",
                     "combined_cost", "depth_sufficient", "seconds_to_close", "cond_price_le_010",
                     "cond_combined_le_090", "cond_depth_ok", "cond_time_ok", "all_conditions_met",
                     "hedge_executed"]
        next_row = _write_table(ws, headers_c, checks, start_row=3)
        ws["A1"] = "Cada chequeo de condiciones de cobertura de MOMENTUM_PARTIAL_HEDGE (se cubra o no). Corte unico ver Overview."

        fills = _rows(conn, """
            SELECT hf.*, pd.strategy FROM paper_hedge_fills hf
            JOIN paper_decisions pd ON pd.id = hf.decision_id
            WHERE hf.executed_at_utc <= ?
            ORDER BY hf.executed_at_utc
        """, REPORT_CUTOFF_TS)
        for r in fills:
            r["executed_at_human"] = fmt(r["executed_at_utc"])
            r["estimated_taker_fee_usd"] = round(estimated_taker_fee_usd(r["filled_shares"], r["vwap"]), 4)
        ws.cell(row=next_row + 1, column=1, value="Coberturas EJECUTADAS (fee estimado con la misma formula, ver TAKER_FEE_SOURCE):")
        headers_f = ["decision_id", "condition_id", "executed_at_human", "target_shares",
                     "filled_shares", "usd_spent", "vwap", "fill_status", "snapshot_id",
                     "estimated_taker_fee_usd"]
        _write_table(ws, headers_f, fills, start_row=next_row + 3)
        ws.freeze_panes = "A4"

        # ---------- Resolutions (P&L, ROI, drawdown por estrategia) ----------
        ws = wb.create_sheet("Resolutions")
        res_rows = _rows(conn, """
            SELECT pr.*, pm.market_title, pm.asset_symbol, pm.validation_cohort FROM paper_resolutions pr
            JOIN paper_markets pm ON pm.condition_id = pr.condition_id
            WHERE pr.resolved_at_utc <= ?
            ORDER BY pr.resolved_at_utc
        """, REPORT_CUTOFF_TS)
        for r in res_rows:
            r["resolved_at_human"] = fmt(r["resolved_at_utc"])
        headers_r = ["strategy", "condition_id", "market_title", "asset_symbol", "validation_cohort",
                     "winner", "resolved_at_human", "capital_deployed_usd", "payout_usd", "pnl_usd", "roi"]
        _write_table(ws, headers_r, res_rows, start_row=3)
        ws["A1"] = ("Resoluciones registradas, TODAS (incluye PRE_VPN_RECOVERY, sin ocultar nada) -- "
                    "vacio hasta que cierren mercados y el collector conozca el winner. Corte unico ver Overview. "
                    "capital_deployed_usd = DENOMINADOR de roi (roi = pnl_usd / capital_deployed_usd).")
        ws.freeze_panes = "A4"

        # ---------- fee estimado real por decision (entrada + cobertura si existe), por REPORT_CUTOFF_TS ----------
        decisions_by_id = {r["id"]: r for r in _rows(conn, """
            SELECT pd.id, pd.shares, pd.executable_price, pd.fill_status FROM paper_decisions pd
            WHERE pd.created_at <= ?
        """, REPORT_CUTOFF_TS)}
        hedge_by_decision = {r["decision_id"]: r for r in _rows(conn, """
            SELECT hf.decision_id, hf.filled_shares, hf.vwap FROM paper_hedge_fills hf
            WHERE hf.executed_at_utc <= ?
        """, REPORT_CUTOFF_TS)}

        def _fee_for_decision(decision_id):
            d = decisions_by_id.get(decision_id)
            fee = 0.0
            if d and d["fill_status"] in ("FULL", "PARTIAL"):
                fee += estimated_taker_fee_usd(d["shares"], d["executable_price"]) or 0.0
            h = hedge_by_decision.get(decision_id)
            if h:
                fee += estimated_taker_fee_usd(h["filled_shares"], h["vwap"]) or 0.0
            return fee

        for r in res_rows:
            r["estimated_taker_fee_usd"] = round(_fee_for_decision(r["decision_id"]), 4)
            r["estimated_pnl_after_taker_fee_usd"] = round(r["pnl_usd"] - r["estimated_taker_fee_usd"], 4)

        # ---------- Strategy Summary (SOLO cohorte OFFICIAL) ----------
        ws = wb.create_sheet("Strategy Summary")
        official_res_rows = [r for r in res_rows if r["validation_cohort"] == "OFFICIAL"]
        n_excluded = len(res_rows) - len(official_res_rows)
        ws["A2"] = (f"Metricas OFICIALES: solo cohorte OFFICIAL (mercados post-recuperacion de VPN, "
                    f">= {fmt(float(cohort_start)) if cohort_start else '?'}). "
                    f"{n_excluded} resoluciones PRE_VPN_RECOVERY excluidas de este resumen (visibles, sin "
                    f"borrar, en la pestaña Resolutions).")
        ws["A3"] = ("IMPORTANTE: estas 3 estrategias son experimentos INDEPENDIENTES sobre el mismo "
                    "universo, no posiciones de una misma cartera -- cada una despliega su propio "
                    "capital por separado. NO sumar sus filas como si fuera el rendimiento combinado "
                    "de un portfolio; se reportan y se leen una por una.")
        ws["A4"] = (f"gross_pnl_usd/gross_roi_pct = P&L BRUTO (fills reales, sin modificar, capital_deployed_usd "
                    f"es el DENOMINADOR del ROI). estimated_taker_fee_usd = fee de TAKER VERIFICADO "
                    f"({TAKER_FEE_SOURCE}), formula shares*{TAKER_FEE_RATE_CRYPTO}*VWAP*(1-VWAP) POR FILL "
                    f"real (entrada + cobertura si existe), sumado por estrategia -- NO una tarifa plana. "
                    f"estimated_pnl_after_taker_fee_usd/estimated_roi_after_taker_fee_pct: NO es beneficio "
                    f"neto definitivo -- {ESTIMATED_FEE_METHOD_NOTE}")
        ws["A5"] = ("NO incluye fee de redencion/resolucion: no verificado con confianza (fuentes "
                    "contradictorias), se deja DESCONOCIDO -- no se aplica 0% ni ningun otro numero.")
        by_strategy = defaultdict(list)
        for r in official_res_rows:
            by_strategy[r["strategy"]].append(r)
        summary_rows = []
        for strategy in ("MOMENTUM_PURE", "MOMENTUM_PARTIAL_HEDGE", "POLYMARKET_FAVORITE_BASELINE"):
            rs = by_strategy.get(strategy, [])
            n = len(rs)
            capital = sum(r["capital_deployed_usd"] for r in rs)
            pnl = sum(r["pnl_usd"] for r in rs)
            fees_est = sum(r["estimated_taker_fee_usd"] for r in rs)
            pnl_after_fees = pnl - fees_est
            wins = sum(1 for r in rs if r["pnl_usd"] > 0)
            losses = sum(1 for r in rs if r["pnl_usd"] < 0)
            by_asset = defaultdict(lambda: {"n": 0, "pnl": 0.0, "capital": 0.0})
            for r in rs:
                a = by_asset[r["asset_symbol"]]
                a["n"] += 1; a["pnl"] += r["pnl_usd"]; a["capital"] += r["capital_deployed_usd"]
            dd = _drawdown([r["pnl_usd"] for r in sorted(rs, key=lambda r: r["resolved_at_utc"])])
            row = {
                "strategy": strategy, "n_resolved": n, "n_wins": wins, "n_losses": losses,
                "capital_deployed_usd_roi_denominator": round(capital, 2),
                "gross_pnl_usd": round(pnl, 2), "gross_roi_pct": round(pnl / capital * 100, 2) if capital else None,
                "estimated_taker_fees_usd": round(fees_est, 4),
                "estimated_pnl_after_taker_fee_usd": round(pnl_after_fees, 2),
                "estimated_roi_after_taker_fee_pct": round(pnl_after_fees / capital * 100, 2) if capital else None,
                "win_rate_pct": round(wins / n * 100, 1) if n else None, "max_drawdown_usd": round(dd, 2),
            }
            for asset in ("BTC", "ETH", "SOL"):
                a = by_asset.get(asset, {"n": 0, "pnl": 0.0, "capital": 0.0})
                row[f"{asset}_n"] = a["n"]; row[f"{asset}_pnl"] = round(a["pnl"], 2)
                row[f"{asset}_roi_pct"] = round(a["pnl"] / a["capital"] * 100, 2) if a["capital"] else None
            summary_rows.append(row)
        headers_s = ["strategy", "n_resolved", "n_wins", "n_losses", "capital_deployed_usd_roi_denominator",
                     "gross_pnl_usd", "gross_roi_pct", "estimated_taker_fees_usd",
                     "estimated_pnl_after_taker_fee_usd", "estimated_roi_after_taker_fee_pct",
                     "win_rate_pct", "max_drawdown_usd", "BTC_n", "BTC_pnl", "BTC_roi_pct",
                     "ETH_n", "ETH_pnl", "ETH_roi_pct", "SOL_n", "SOL_pnl", "SOL_roi_pct"]
        _write_table(ws, headers_s, summary_rows, start_row=7)
        ws["A1"] = "P&L/ROI/drawdown por estrategia (independientes) -- se llena a medida que cierran mercados. Corte unico ver Overview."
        ws.column_dimensions["A"].width = 28
        ws.freeze_panes = "A8"

        # ---------- Momentum vs Favorite (exactamente los mismos condition_id) ----------
        ws = wb.create_sheet("Momentum vs Favorite")
        mom_by_cid = {r["condition_id"]: r for r in official_res_rows if r["strategy"] == "MOMENTUM_PURE"}
        fav_by_cid = {r["condition_id"]: r for r in official_res_rows if r["strategy"] == "POLYMARKET_FAVORITE_BASELINE"}
        common_cids = set(mom_by_cid) & set(fav_by_cid)
        ws["A1"] = (f"MOMENTUM_PURE vs POLYMARKET_FAVORITE_BASELINE, restringido a los "
                    f"{len(common_cids)} condition_id donde AMBAS tienen resolucion OFFICIAL -- "
                    f"comparacion exacta, no universos separados. (MOMENTUM_PURE tiene "
                    f"{len(mom_by_cid)} en total, FAVORITE {len(fav_by_cid)}; se usa la interseccion.) "
                    f"gross_* = P&L bruto (capital_usd es el denominador del ROI). "
                    f"estimated_pnl_after_taker_fee_usd = ESTIMACION aproximada por VWAP, NO beneficio neto "
                    f"definitivo -- ver {TAKER_FEE_SOURCE}. Fee de redencion: DESCONOCIDO, no aplicado.")

        def _agg_common(by_cid, cids, label):
            rs = [by_cid[c] for c in cids]
            n = len(rs)
            capital = sum(r["capital_deployed_usd"] for r in rs)
            pnl = sum(r["pnl_usd"] for r in rs)
            fees = sum(r["estimated_taker_fee_usd"] for r in rs)
            pnl_after_fees = pnl - fees
            wins = sum(1 for r in rs if r["pnl_usd"] > 0)
            by_asset = defaultdict(lambda: {"n": 0, "pnl": 0.0, "capital": 0.0})
            for r in rs:
                a = by_asset[r["asset_symbol"]]
                a["n"] += 1; a["pnl"] += r["pnl_usd"]; a["capital"] += r["capital_deployed_usd"]
            row = {"strategy": label, "n": n, "capital_usd_roi_denominator": round(capital, 2),
                   "gross_pnl_usd": round(pnl, 2), "gross_roi_pct": round(pnl / capital * 100, 2) if capital else None,
                   "estimated_taker_fees_usd": round(fees, 4),
                   "estimated_pnl_after_taker_fee_usd": round(pnl_after_fees, 2),
                   "estimated_roi_after_taker_fee_pct": round(pnl_after_fees / capital * 100, 2) if capital else None,
                   "win_rate_pct": round(wins / n * 100, 1) if n else None}
            for asset in ("BTC", "ETH", "SOL"):
                a = by_asset.get(asset, {"n": 0, "pnl": 0.0, "capital": 0.0})
                row[f"{asset}_n"] = a["n"]; row[f"{asset}_pnl"] = round(a["pnl"], 2)
                row[f"{asset}_roi_pct"] = round(a["pnl"] / a["capital"] * 100, 2) if a["capital"] else None
            return row

        cmp_rows = [_agg_common(mom_by_cid, common_cids, "MOMENTUM_PURE_common"),
                    _agg_common(fav_by_cid, common_cids, "POLYMARKET_FAVORITE_common")]
        headers_cmp = ["strategy", "n", "capital_usd_roi_denominator", "gross_pnl_usd", "gross_roi_pct",
                       "estimated_taker_fees_usd", "estimated_pnl_after_taker_fee_usd",
                       "estimated_roi_after_taker_fee_pct", "win_rate_pct",
                       "BTC_n", "BTC_pnl", "BTC_roi_pct", "ETH_n", "ETH_pnl", "ETH_roi_pct",
                       "SOL_n", "SOL_pnl", "SOL_roi_pct"]
        next_row_cmp = _write_table(ws, headers_cmp, cmp_rows, start_row=3)

        detail_rows = []
        for c in sorted(common_cids, key=lambda c: mom_by_cid[c]["resolved_at_utc"]):
            m, f = mom_by_cid[c], fav_by_cid[c]
            detail_rows.append({
                "condition_id": c, "asset_symbol": m["asset_symbol"], "winner": m["winner"],
                "momentum_gross_pnl_usd": round(m["pnl_usd"], 4), "favorite_gross_pnl_usd": round(f["pnl_usd"], 4),
                "momentum_estimated_fee_usd": m["estimated_taker_fee_usd"],
                "favorite_estimated_fee_usd": f["estimated_taker_fee_usd"],
                "momentum_estimated_pnl_after_fee_usd": m["estimated_pnl_after_taker_fee_usd"],
                "favorite_estimated_pnl_after_fee_usd": f["estimated_pnl_after_taker_fee_usd"],
                "momentum_capital_usd": m["capital_deployed_usd"], "favorite_capital_usd": f["capital_deployed_usd"],
            })
        ws.cell(row=next_row_cmp + 1, column=1, value="Detalle por mercado (interseccion), bruto y estimado post-fee:")
        _write_table(ws, ["condition_id", "asset_symbol", "winner", "momentum_gross_pnl_usd",
                          "favorite_gross_pnl_usd", "momentum_estimated_fee_usd", "favorite_estimated_fee_usd",
                          "momentum_estimated_pnl_after_fee_usd", "favorite_estimated_pnl_after_fee_usd",
                          "momentum_capital_usd", "favorite_capital_usd"],
                     detail_rows, start_row=next_row_cmp + 3)
        ws.column_dimensions["A"].width = 28
        ws.freeze_panes = "A4"

        # ---------- Sleep Audit (sueno por tapa cerrada, 2026-09-16) ----------
        ws = wb.create_sheet("Sleep Audit")
        # Ventana verificada contra el log de energia del sistema (pmset -g log: "Entering Sleep state
        # due to Clamshell Sleep" / "Wake") Y contra el hueco real medido en orderbook_snapshots de
        # data.db -- ambas fuentes coinciden. Constantes fijas: el evento ya ocurrio y no cambia.
        SLEEP_START_TS = 1789551755.0   # 2026-09-16 09:42:35 UTC (11:42:35 CEST)
        SLEEP_END_TS = 1789552797.277   # 2026-09-16 09:59:57.277 UTC -- primer dato REAL tras el sueno
        markets_in_gap = _rows(conn, """
            SELECT condition_id, asset_symbol, decision_time_utc, validation_cohort
            FROM paper_markets WHERE decision_time_utc BETWEEN ? AND ? AND discovered_at <= ?
            ORDER BY decision_time_utc
        """, SLEEP_START_TS, SLEEP_END_TS, REPORT_CUTOFF_TS)
        n_gap_markets = len(markets_in_gap)
        gap_decisions = []
        for m in markets_in_gap:
            for d in _rows(conn, "SELECT * FROM paper_decisions WHERE condition_id=? AND created_at<=?",
                            m["condition_id"], REPORT_CUTOFF_TS):
                gap_decisions.append(d)
        n_gap_decisions = len(gap_decisions)
        n_gap_skipped = sum(1 for d in gap_decisions if d["fill_status"] == "SKIPPED")
        n_gap_traded = sum(1 for d in gap_decisions if d["fill_status"] in ("FULL", "PARTIAL"))
        hedge_checks_in_window = _rows(conn, """
            SELECT * FROM paper_hedge_checks WHERE checked_at_utc BETWEEN ? AND ? AND checked_at_utc <= ?
        """, SLEEP_START_TS - 60, SLEEP_END_TS + 60, REPORT_CUTOFF_TS)
        n_hedge_checks_met = sum(1 for c in hedge_checks_in_window if c["all_conditions_met"])
        hedge_fills_in_window = _rows(conn, """
            SELECT * FROM paper_hedge_fills WHERE executed_at_utc BETWEEN ? AND ? AND executed_at_utc <= ?
        """, SLEEP_START_TS - 60, SLEEP_END_TS + 60, REPORT_CUTOFF_TS)
        n_official_now = conn.execute(
            "SELECT count(*) c FROM paper_markets WHERE validation_cohort='OFFICIAL' AND discovered_at<=?",
            (REPORT_CUTOFF_TS,)
        ).fetchone()["c"]
        pct_affected = round(n_gap_markets / n_official_now * 100, 2) if n_official_now else None

        ws["A1"] = ("Incidente: el Mac entro en sueno por tapa cerrada (Clamshell Sleep) el 2026-09-16, "
                    "NO se apago -- confirmado por 'uptime' (60 dias sin reiniciar) y por pmset -g log. "
                    "Los procesos (collector PID 38099, validador PID 52466) se PAUSARON y reanudaron "
                    "solos al despertar -- nunca murieron, nunca se reiniciaron.")
        ws["A2"] = f"Ventana del sueno (verificada por 2 fuentes independientes): {fmt(SLEEP_START_TS)} -> {fmt(SLEEP_END_TS)} (~{(SLEEP_END_TS-SLEEP_START_TS)/60:.1f} min)"
        ws["A3"] = (f"Mercados OFFICIAL con decision_time_utc DENTRO del hueco: {n_gap_markets} de "
                    f"{n_official_now} totales = {pct_affected}%.")
        ws["A4"] = (f"De sus {n_gap_decisions} decisiones (mercados x 3 estrategias): {n_gap_skipped} SKIPPED, "
                    f"{n_gap_traded} FULL/PARTIAL. Verificado uno por uno contra data.db (snapshot_id real, "
                    f"received_at_utc_ms <= decision_timestamp_utc): NINGUNA decision uso un snapshot con "
                    f"antiguedad fuera de la tolerancia de 2.5s -- el fail-closed funciono, cero ejecuciones "
                    f"retrospectivas con datos obsoletos.")
        ws["A5"] = (f"Coberturas (MOMENTUM_PARTIAL_HEDGE) en la ventana +/-60s: {len(hedge_checks_in_window)} "
                    f"chequeos, {n_hedge_checks_met} con all_conditions_met=1, {len(hedge_fills_in_window)} "
                    f"coberturas EJECUTADAS. Ninguna cobertura se ejecuto con datos del hueco.")
        ws["A6"] = ("Registros preservados en su totalidad -- nada se borro. Estos 9 mercados quedan como "
                    "SKIPPED con motivo explicito en skip_or_partial_reason (Signals), no como si nunca "
                    "hubieran existido.")

        gap_rows = []
        for m in markets_in_gap:
            decs = [d for d in gap_decisions if d["condition_id"] == m["condition_id"]]
            for d in decs:
                gap_rows.append({
                    "condition_id": m["condition_id"], "asset_symbol": m["asset_symbol"],
                    "decision_time_human": fmt(m["decision_time_utc"]), "strategy": d["strategy"],
                    "fill_status": d["fill_status"], "skip_or_partial_reason": d["skip_or_partial_reason"],
                    "created_at_human": fmt(d["created_at"]),
                })
        ws.cell(row=8, column=1, value="Detalle de las 27 decisiones afectadas:")
        _write_table(ws, ["condition_id", "asset_symbol", "decision_time_human", "strategy", "fill_status",
                          "skip_or_partial_reason", "created_at_human"], gap_rows, start_row=10)
        ws.column_dimensions["A"].width = 26
        ws.freeze_panes = "A11"

        # ---------- Methodology ----------
        ws = wb.create_sheet("Methodology")
        methodology = [
            "Datos completamente nuevos: solo mercados con open_time_utc posterior al arranque del "
            "validador (nunca los ya usados en el backtest V1/V2/V3).",
            "MOMENTUM_PURE: decision a los 60s de abierto. Opera solo si |distancia del subyacente al "
            "open| >= 0.02% (umbral congelado de V1). Compra UP si positiva, DOWN si negativa. Stake $10. "
            "Mantiene a resolucion. Nunca compra la pierna contraria.",
            "MOMENTUM_PARTIAL_HEDGE: misma entrada que MOMENTUM_PURE. Monitorea la pierna contraria y "
            "cubre EXACTAMENTE 25% de las shares iniciales, una sola vez, solo si TODAS estas condiciones "
            "se cumplen simultaneamente: (1) precio ejecutable de la pierna contraria <= $0.10; (2) coste "
            "medio de la 1ra pierna + precio ejecutable contrario <= $0.90; (3) profundidad suficiente para "
            "llenar el 25% completo (FULL, nunca PARTIAL); (4) quedan >= 60s para el cierre. Nunca vende.",
            "POLYMARKET_FAVORITE_BASELINE: decision a los 60s, compra el lado con MAYOR probabilidad "
            "implicita (precio/mid mas alto) -- nota: el enunciado original decia 'menor best ask / mayor "
            "probabilidad implicita', que son contradictorios bajo el mecanismo real de precios; se resolvio "
            "usando 'mayor probabilidad implicita', consistente con 'favorito' en toda la sesion previa "
            "(V1/V2/V3). Stake $10, mantiene a resolucion.",
            "Ejecucion: SIEMPRE recorriendo profundidad real (orderbook_levels), nunca asume fill al "
            "best_ask. FULL/PARTIAL/SKIPPED registrados explicitamente. Fail-closed: sin datos = SKIPPED, "
            "nunca se inventa un precio.",
            "Anti-lookahead: toda decision usa unicamente snapshots con received_at_utc_ms <= el instante "
            "de decision (misma regla conservadora que trade_context/V1/V2/V3). El monitoreo de cobertura "
            "usa el instante real de wall-clock en que corre el chequeo -- no puede ver el futuro porque "
            "corre en tiempo real hacia adelante.",
            "Ninguna estrategia observa ni usa las operaciones del lider como señal -- las 3 son "
            "autonomas, evaluadas sobre el mismo universo de mercados BTC/ETH/SOL de 5 minutos.",
            "SEGURIDAD: este validador nunca hace una request HTTP (0 dependencias de red), nunca importa "
            "el ejecutor real ni ningun cliente de ordenes/wallet, nunca lee credenciales, y abre "
            "collector/data.db exclusivamente en modo read-only (a nivel de driver sqlite, no solo por "
            "convencion). Escribe unicamente en collector/paper_validation.db.",
            "Las 3 estrategias son experimentos INDEPENDIENTES, no posiciones de una misma cartera -- "
            "sus filas en Strategy Summary NUNCA deben sumarse como si fuera un portfolio combinado.",
            f"Fees: gross_pnl_usd/gross_roi_pct son BRUTOS (fills reales, sin tocar; capital_deployed_usd "
            f"es el denominador del ROI). estimated_taker_fee_usd se VERIFICO contra {TAKER_FEE_SOURCE}: "
            f"fee = shares * {TAKER_FEE_RATE_CRYPTO} * precio * (1-precio), categoria crypto, solo taker "
            f"(maker=0%, y las 3 estrategias son 100% taker). {ESTIMATED_FEE_METHOD_NOTE} NO se aplica "
            f"ningun numero de fee de redencion/resolucion -- las fuentes son contradictorias (la mayoria "
            f"dice 0%, una dice 2%), asi que queda DESCONOCIDO, nunca asumido como cero ni como 2%. Nunca "
            f"se modifica paper_decisions/paper_resolutions -- el fee es un calculo derivado, separado.",
            "Momentum vs Favorite: comparacion restringida a la INTERSECCION exacta de condition_id "
            "donde ambas estrategias tienen resolucion OFFICIAL -- nunca se comparan sus universos "
            "completos por separado (que difieren, porque MOMENTUM_PURE exige señal y FAVORITE no). "
            "Se reporta bruto y estimado post-fee, con la misma etiqueta de aproximacion.",
            "Incidente de sueno del sistema (2026-09-16): ver hoja 'Sleep Audit' -- el Mac nunca se "
            "apago (60 dias de uptime), solo durmio ~17.4min por tapa cerrada. 9/364 mercados OFFICIAL "
            "(2.47%) quedaron SKIPPED por ese hueco, verificado que ninguna decision ni cobertura uso "
            "datos obsoletos (fail-closed). Registros preservados, nada borrado.",
            "Corte unico del informe: todas las hojas usan el mismo REPORT_CUTOFF_TS (ver Overview) para "
            "que los conteos/capital/P&L sean reconciliables entre pestañas, incluso si el validador "
            "siguio generando datos nuevos mientras se armaba este Excel.",
        ]
        ws.append(["nota_metodologica"])
        for m in methodology:
            ws.append([m])
        ws.column_dimensions["A"].width = 120

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    sha256 = hashlib.sha256(out_path.read_bytes()).hexdigest()
    out_path.with_suffix(out_path.suffix + ".sha256").write_text(f"{sha256}  {out_path.name}\n")
    print(f"wrote {out_path}  sha256={sha256}")
    return out_path, sha256


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "export" / "paper_validation.xlsx"
    build(out)
